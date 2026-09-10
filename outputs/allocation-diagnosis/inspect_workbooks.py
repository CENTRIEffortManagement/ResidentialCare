"""Read saved workbook tables and embedded M; never write source workbooks."""
from pathlib import Path
import base64, collections, io, json, re, struct, zipfile
import xml.etree.ElementTree as ET
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / 'CLIENT/DATExx-Whiddon/UNITS/Unit1/2. Calculations'
OUT = Path(__file__).resolve().parent
N = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def read_book(name):
    path = BASE / name
    tables = {}
    with zipfile.ZipFile(path) as z:
        for item in z.namelist():
            if item.startswith('xl/tables/') and item.endswith('.xml'):
                root = ET.fromstring(z.read(item))
                tables[root.attrib['name']] = root.attrib['ref']
        for item in z.namelist():
            if item.startswith('customXml/') and item.endswith('.xml'):
                root = ET.fromstring(z.read(item))
                if root.tag.endswith('DataMashup') and root.text:
                    raw = base64.b64decode(root.text)
                    size = struct.unpack_from('<I', raw, 4)[0]
                    with zipfile.ZipFile(io.BytesIO(raw[8:8+size])) as inner:
                        for member in inner.namelist():
                            if member.endswith('.m'):
                                code = inner.read(member).decode('utf-8-sig')
                                (OUT / (name + '.embedded.m')).write_text(code, encoding='utf-8')
    if name == 'Shifts.xlsx':
        print(json.dumps({'file': name, 'embedded_query_only': True, 'tables': tables}))
        return
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True, keep_links=False)
    data = {}
    for ws in wb:
        for tab in ws.tables.values():
            rows = list(ws[tab.ref])
            headers = [c.value for c in rows[0]]
            data[tab.name] = [dict(zip(headers, [c.value for c in row])) for row in rows[1:]]
    summary = {'file': name, 'modified': path.stat().st_mtime, 'tables': {k:{'rows':len(v), 'columns':list(v[0]) if v else []} for k,v in data.items()}}
    (OUT / (name + '.tables.json')).write_text(json.dumps(data, default=str), encoding='utf-8')
    print(json.dumps(summary))
    wb.close()

if __name__ == '__main__':
    import sys
    for name in sys.argv[1:]:
        read_book(name)
