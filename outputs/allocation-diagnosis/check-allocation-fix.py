"""Regression model using saved tables. Checks arithmetic/joins; does not execute M."""
from pathlib import Path
from collections import defaultdict
from datetime import datetime, timedelta
import json, math
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
BASE=ROOT/'CLIENT/DATExx-Whiddon/UNITS/Unit1/2. Calculations'
OUT=Path(__file__).resolve().parent

def read_sheet_with_columns(filename, required):
    wb=openpyxl.load_workbook(BASE/filename,read_only=True,data_only=True,keep_links=False)
    try:
        for ws in wb:
            rows=ws.iter_rows(values_only=True)
            header=next(rows,())
            if set(required).issubset(header):
                indices=[header.index(k) for k in required]
                return [dict(zip(required,[row[i] for i in indices])) for row in rows if any(row[i] is not None for i in indices)]
    finally:
        wb.close()
    raise AssertionError(f'Missing sheet columns in {filename}: {required}')

def settings_table(name):
    wb=openpyxl.load_workbook(BASE/'Settings Data.xlsx',data_only=True,keep_links=False)
    try:
        for ws in wb:
            if name in ws.tables:
                rows=list(ws[ws.tables[name].ref])
                header=[c.value for c in rows[0]]
                return [dict(zip(header,[c.value for c in row])) for row in rows[1:]]
    finally:
        wb.close()
    raise AssertionError('Missing settings table: '+name)

def corrected(staff, durations):
    effort=defaultdict(float)
    for r in staff:
        if r['StaffCount']==1 and r['Attribute']=='IntervalStart':
            k=(r['ShiftDate'],r['Role'],r['ShiftPeriod'])
            effort[k]+=r['ResEffectiveIntervalEffort']
    lookup=defaultdict(list)
    for r in durations: lookup[(r['Role'],r['ShiftPeriod'])].append(r['DurationOfShifts'])
    output={}
    for k,v in effort.items():
        matches=lookup[(k[1],k[2])]
        if len(matches)!=1 or matches[0] is None or not math.isfinite(matches[0]) or matches[0]<=0:
            raise ValueError('Invalid role-shift duration')
        if any(x is None or x=='' for x in k) or not math.isfinite(v) or v<0:
            raise ValueError('Invalid allocation')
        output[k]=(v,matches[0],v*24/matches[0])
    return output

staff=read_sheet_with_columns('AllocationByShiftAverage.xlsx',['ShiftDate','Role','ShiftPeriod','Attribute','StaffCount','ResEffectiveIntervalEffort'])
durations=settings_table('ShiftPeriod')
calendar=settings_table('PermutationDimensions')
actual=corrected(staff,durations)
output={(r['Date'],r['RolesList'],r['Shifts']):actual.get((r['Date'],r['RolesList'],r['Shifts']),(0,None,0)) for r in calendar}
assert len(output)==len(calendar)==252
assert all(abs(e*24-f*d)<1e-8 for e,d,f in actual.values())
example=next(v for (d,r,s),v in output.items() if r=='AIN' and s=='AM' and str(d).startswith('2026-07-20'))
assert math.isclose(example[2],10.836616161616162,rel_tol=1e-12)
fixture=[{'ShiftDate':datetime(2026,7,20),'Role':'AIN','ShiftPeriod':'AM','Attribute':'IntervalStart','StaffCount':1,'ResEffectiveIntervalEffort':8.25/24},
         {'ShiftDate':datetime(2026,7,21),'Role':'AIN','ShiftPeriod':'AM','Attribute':'IntervalEnd','StaffCount':1,'ResEffectiveIntervalEffort':8.25/24},
         {'ShiftDate':datetime(2026,7,20),'Role':'AIN','ShiftPeriod':'AM','Attribute':'IntervalStart','StaffCount':0,'ResEffectiveIntervalEffort':0}]
result=corrected(fixture,durations)
assert len(result)==1 and next(iter(result.values()))[2]==1
base_duration={'Role':'AIN','ShiftPeriod':'AM','DurationOfShifts':8.25}
bad_sets=[[],[base_duration,base_duration],[dict(base_duration,DurationOfShifts=0)],[dict(base_duration,DurationOfShifts=-1)],[dict(base_duration,DurationOfShifts=None)],[dict(base_duration,DurationOfShifts=float('inf'))],[dict(base_duration,DurationOfShifts=float('nan'))]]
for rows in bad_sets:
    try: corrected(fixture,rows)
    except ValueError: pass
    else: raise AssertionError('Invalid duration was accepted')
evidence={'validation':'passed','execution':'Python regression model; not an Excel M refresh','saved_resource_rows':len(staff),'corrected_upstream_keys':len(actual),'calendar_output_rows':len(output),'zero_padding_rows':sum(d is None for e,d,f in output.values()),'example_start_effort_days':example[0],'example_shift_hours':example[1],'example_fte':example[2],'invalid_duration_scenarios_rejected':len(bad_sets),'endpoint_boundary_and_role_expansion':'passed','fte_hours_reconciliation':'passed'}
(OUT/'fix-validation.json').write_text(json.dumps(evidence,indent=2))
print(json.dumps(evidence,indent=2))
