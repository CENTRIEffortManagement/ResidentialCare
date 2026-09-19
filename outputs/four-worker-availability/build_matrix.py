"""Read the pre-repair saved Unit1 tables and explain each sampled worker shift.

This is a read-only historical diagnostic. It does not refresh or modify any workbook.
It models the saved pre-repair rule; build_expected_matrix.ps1 models the revised rule.
"""

from __future__ import annotations

import csv
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


ROOT = Path(__file__).resolve().parents[2]
UNIT = ROOT / "CLIENT" / "DATExx-Whiddon" / "UNITS" / "Unit1"
AVAILABILITY = UNIT / "2. Calculations" / "Capacity-ShiftAvailability.xlsx"
SETTINGS = UNIT / "2. Calculations" / "Settings Data.xlsx"
OUTPUT = Path(__file__).resolve().parent


def records(sheet):
    header, *rows = sheet.values
    return [dict(zip(header, row)) for row in rows if any(value is not None for value in row)]


def union(intervals):
    merged = []
    for start, end in sorted(intervals):
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((start, end))
    return merged


def clip(intervals, start, end):
    return [(max(a, start), min(b, end)) for a, b in intervals if a < end and b > start]


def subtract(base, cuts):
    parts = list(base)
    for start, end in cuts:
        next_parts = []
        for a, b in parts:
            if end <= a or start >= b:
                next_parts.append((a, b))
            else:
                if start > a:
                    next_parts.append((a, start))
                if end < b:
                    next_parts.append((end, b))
        parts = next_parts
    return parts


def hours(intervals):
    return sum((end - start).total_seconds() / 3600 for start, end in intervals)


def stamp(value):
    return value.strftime("%Y-%m-%d %H:%M")


def windows(intervals):
    return "; ".join(f"{stamp(start)} to {stamp(end)}" for start, end in intervals)


def source_text(items):
    return "; ".join(
        f"#{item['SourceRecord']} {stamp(item['Start'])} to {stamp(item['End'])}"
        for item in sorted(items, key=lambda x: (x["Start"], x["SourceRecord"]))
    )


availability_book = load_workbook(AVAILABILITY, data_only=True, read_only=True)
settings_book = load_workbook(SETTINGS, data_only=True, read_only=False)
shift_sheet = settings_book["ShiftHours"]
shift_duration_cells = shift_sheet[shift_sheet.tables["ShiftDuration"].ref]
shift_header = [cell.value for cell in shift_duration_cells[0]]
shift_values = [cell.value for cell in shift_duration_cells[1]]
allowance = float(dict(zip(shift_header, shift_values))["ShiftDuration"])

workers = {str(row["EmployeeID"]): row for row in records(availability_book["ReconciledWorkers_Eligible"])}
raw_by_record = {
    int(row["SourceRecord"]): row
    for row in records(availability_book["IMPORT Availability Leave Sourc"])
}
source_by_worker = defaultdict(list)
for row in records(availability_book["AvailabilityRecords"]):
    if row["Issue"] is not None:
        continue
    raw = raw_by_record[int(row["SourceRecord"])]
    reason = raw["Availability or Leave Reason"]
    kind = "AVAIL" if row["Available"] else "UNAVAIL" if "UNAVAIL" in reason else "Leave/other exclusion"
    source_by_worker[str(row["Payroll Code"])].append(
        {
            "SourceRecord": int(row["SourceRecord"]),
            "Start": from_excel(row["Start"]),
            "End": from_excel(row["End"]),
            "Kind": kind,
            "Reason": reason,
        }
    )

saved_shifts = records(availability_book["WorkerShiftSegments"])
published_keys = {
    (str(row["ID"]), from_excel(row["Date"]).date(), row["Shift"])
    for row in records(availability_book["ResDayShift"])
}

daily_cache = {}


def day_evidence(worker_id, day):
    key = worker_id, day
    if key in daily_cache:
        return daily_cache[key]
    start = datetime.combine(day, datetime.min.time())
    end = start + timedelta(days=1)
    entries = [
        row for row in source_by_worker[worker_id]
        if row["Start"] < end and row["End"] > start
    ]
    raw_exclusions = union(
        clip([(row["Start"], row["End"]) for row in entries if row["Kind"] != "AVAIL"], start, end)
    )
    distinct_hours = hours(raw_exclusions)
    whole_day = distinct_hours >= allowance
    effective_exclusions = [(start, end)] if whole_day else raw_exclusions
    result = entries, distinct_hours, whole_day, effective_exclusions
    daily_cache[key] = result
    return result


matrix = []
for saved in saved_shifts:
    worker_id = str(saved["EmployeeID"])
    start = from_excel(saved["ShiftStart"])
    end = from_excel(saved["ShiftEnd"])
    current = datetime.combine(start.date(), datetime.min.time())
    touched_days = []
    while current < end:
        touched_days.append(current.date())
        current += timedelta(days=1)
    touched_entries = {}
    daily = []
    exclusions = []
    for day in touched_days:
        entries, excluded_hours, whole_day, day_exclusions = day_evidence(worker_id, day)
        touched_entries.update((entry["SourceRecord"], entry) for entry in entries)
        daily.append(f"{day}: {excluded_hours:g}h" + (" -> full day" if whole_day else " -> timed only"))
        exclusions.extend(day_exclusions)
    source = list(touched_entries.values())
    has_avail = any(row["Kind"] == "AVAIL" for row in source_by_worker[worker_id])
    if has_avail:
        base = union(clip(
            [(row["Start"], row["End"]) for row in source_by_worker[worker_id] if row["Kind"] == "AVAIL"],
            start, end,
        ))
    else:
        base = [(start, end)]
    shift_cuts = union(clip(exclusions, start, end))
    remaining = subtract(base, shift_cuts)
    calculated_hours = hours(remaining)
    full_shift = len(remaining) == 1 and remaining[0] == (start, end)
    effective_hours = allowance if full_shift else min(calculated_hours, allowance)
    key = worker_id, from_excel(saved["Date"]).date(), saved["Shift"]
    if effective_hours > 0:
        decision = "Positive hours: expected in ResDayShift"
    elif not base:
        decision = "No AVAIL window intersects shift"
    elif any(day_evidence(worker_id, day)[2] for day in touched_days):
        decision = "Whole-day exclusion removes shift"
    else:
        decision = "Timed exclusions cover shift"
    matrix.append(
        {
            "EmployeeID": worker_id,
            "Name": workers[worker_id]["Name"],
            "Role": saved["Role"],
            "Date": key[1].isoformat(),
            "Day": saved["Day"],
            "Shift": saved["Shift"],
            "ShiftStart": stamp(start),
            "ShiftEnd": stamp(end),
            "HasAVAILAnywhereInSource": has_avail,
            "AVAILSourceRecords": source_text([item for item in source if item["Kind"] == "AVAIL"]),
            "UNAVAILSourceRecords": source_text([item for item in source if item["Kind"] == "UNAVAIL"]),
            "LeaveOtherSourceRecords": source_text([item for item in source if item["Kind"] == "Leave/other exclusion"]),
            "DistinctDailyExcludedHoursAndRule": "; ".join(daily),
            "BaselineWithinShift": windows(base),
            "ExclusionsAppliedToShift": windows(shift_cuts),
            "RemainingAvailableSegments": windows(remaining),
            "CalculatedAvailableHours": round(calculated_hours, 8),
            "SavedWorkerShiftSegmentsHours": saved["AvailableHours"],
            "FullShift": full_shift,
            "CalculatedEffectiveShiftHrs": round(effective_hours, 8),
            "PresentInSavedResDayShift": key in published_keys,
            "Decision": decision,
        }
    )

matrix.sort(key=lambda row: (row["EmployeeID"], row["Date"], row["ShiftStart"]))
assert len(matrix) == len(saved_shifts) == 336
assert all(abs(row["CalculatedAvailableHours"] - row["SavedWorkerShiftSegmentsHours"]) < 1e-7 for row in matrix)
assert all((row["CalculatedEffectiveShiftHrs"] > 0) == row["PresentInSavedResDayShift"] for row in matrix)

with (OUTPUT / "day_shift_matrix_saved.csv").open("w", newline="", encoding="utf-8-sig") as handle:
    writer = csv.DictWriter(handle, fieldnames=list(matrix[0]))
    writer.writeheader()
    writer.writerows(matrix)

print("Workers:", ", ".join(sorted(workers)))
print("ShiftDuration:", allowance)
print("Rows:", len(matrix))
print("Saved versus calculated hours: all match")
print("Published rows:", len(published_keys))
print("Decisions:", dict(Counter(row["Decision"] for row in matrix)))
print("Wrote:", OUTPUT / "day_shift_matrix_saved.csv")
